"""
Enterprise Finance Reconciliation Tool
======================================

Purpose:
--------
Performs automated, deterministic, rule-driven financial reconciliation between an
Enterprise System of Record (ERP) and an External Source (Bank / Payment Gateway / Partner).

Key Principles:
---------------
1. Explicit Configuration & Mandatory Schema: Zero substring-based field type inference. Rules (money,
   text, date, required, currency_field) are explicitly declared, validated, and strictly enforced.
2. Raw Identifier Preservation (C1): Ingests CSV (dtype=str) and XLSX (dtype=object) preserving raw
   leading zero string keys ("00123") and canonicalizing float IDs (1001.0 vs "1001") via _recon_key.
3. Structured Money Parsing & Currency Conflict Check (C3): Parses amount, symbol, and embedded currency.
   Validates embedded currency against the record's currency field (CURRENCY_AMOUNT_CONFLICT).
4. Strict Required Field Semantics (C2 & C6): Enforces normalize_text(). Missing required fields on either
   or both sides produce data quality exception status MISSING_REQUIRED_VALUE.
5. Dual Status Architecture & Source Validation (C8): Separates reconciliation_status (MATCH, MISMATCH,
   MISSING_IN_EXTERNAL, MISSING_IN_ERP) from data_quality_status (VALID, INVALID_FIELD_VALUE,
   MISSING_REQUIRED_VALUE, CURRENCY_AMOUNT_CONFLICT, DUPLICATE_KEY_*).
6. Separated Output Grains (C4): Outputs distinct sheets for reconciliation_units (1 row per logical
   reconciliation unit) and source_exceptions (1 row per problematic physical source record).
7. Disambiguated KPIs & Dynamic Exposure (C5 & C7): Reports separate Source Row Exception Rate (%) and
   Reconciliation Unit Exception Rate (%), and calculates dynamic exposure per money field and currency.
8. Auditable Reporting: Multi-sheet Excel workbook with UTC timestamps, SHA-256 hashes, source row lineage,
   and Excel formula injection hardening.
"""

import argparse
import hashlib
import json
import logging
import os
import re
import sys
import uuid
from datetime import date, datetime, timezone
from decimal import Decimal, InvalidOperation
from importlib.metadata import PackageNotFoundError, version
from typing import Any, Dict, List, Optional, Tuple, Union

import openpyxl
import pandas as pd
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

logger = logging.getLogger("reconciliation")

SUPPORTED_EXTENSIONS = {".csv", ".xlsx"}
DEFAULT_TOLERANCE = Decimal("0.01")
DEFAULT_KEY_COLUMN = "invoice_id"
INTERNAL_PREFIX = "_recon_"
RECON_KEY = f"{INTERNAL_PREFIX}key"

# Machine-readable reconciliation status constants (Unit-level cardinality)
STATUS_MATCH = "MATCH"
STATUS_MISMATCH = "MISMATCH"
STATUS_MISSING_IN_EXTERNAL = "MISSING_IN_EXTERNAL"
STATUS_MISSING_IN_ERP = "MISSING_IN_ERP"

# Machine-readable data quality status constants (Data Quality dimension)
DQ_VALID = "VALID"
DQ_MISSING_PRIMARY_KEY = "MISSING_PRIMARY_KEY"
DQ_DUPLICATE_IN_ERP = "DUPLICATE_KEY_IN_ERP"
DQ_DUPLICATE_IN_EXTERNAL = "DUPLICATE_KEY_IN_EXTERNAL"
DQ_DUPLICATE_IN_BOTH = "DUPLICATE_KEY_IN_BOTH"
DQ_INVALID_FIELD_VALUE = "INVALID_FIELD_VALUE"
DQ_MISSING_REQUIRED_VALUE = "MISSING_REQUIRED_VALUE"
DQ_CURRENCY_AMOUNT_CONFLICT = "CURRENCY_AMOUNT_CONFLICT"

INVALID_VALUE_SENTINEL = "<INVALID_FIELD_VALUE>"

# Allowed currency codes and symbols for money parsing
CURRENCY_PATTERN = re.compile(
    r"[$₹€£¥]|Rs\.?|\b(USD|EUR|INR|GBP|CAD|AUD|JPY|CHF|CNY|HKD|NZD)\b",
    re.IGNORECASE,
)

SYMBOL_TO_CURRENCY = {
    "$": "USD",
    "€": "EUR",
    "₹": "INR",
    "£": "GBP",
    "¥": "JPY",
    "Rs": "INR",
    "Rs.": "INR",
}


class ParsedMoney:
    """Structured monetary parsing result (C3)."""

    def __init__(
        self,
        amount: Optional[Decimal],
        explicit_currency: Optional[str],
        symbol: Optional[str],
        valid: bool,
    ):
        self.amount = amount
        self.explicit_currency = explicit_currency
        self.symbol = symbol
        self.valid = valid


class ReconciliationError(Exception):
    """Base exception for reconciliation failures."""

    pass


class ConfigurationError(ReconciliationError):
    """Raised when configuration parameters or field rules are invalid."""

    pass


class SchemaError(ReconciliationError):
    """Raised when dataset schemas or header validations fail."""

    pass


class IngestionError(ReconciliationError):
    """Raised when file reading or ingestion fails."""

    pass


class OutputError(ReconciliationError):
    """Raised when output file writing or overwrite checks fail."""

    pass


def get_tool_version() -> str:
    """Returns dynamic package version or fallback."""
    try:
        return version("enterprise-finance-reconciliation-tool")
    except PackageNotFoundError:
        return "1.0.0"
    except Exception:
        return "1.0.0"


def calculate_file_hash(file_path: str) -> str:
    """Calculates SHA-256 hash of a file for auditability."""
    sha256 = hashlib.sha256()
    with open(file_path, "rb") as f:
        for chunk in iter(lambda: f.read(65536), b""):
            sha256.update(chunk)
    return sha256.hexdigest()


def normalize_text(val: Any) -> Optional[str]:
    """
    Standard text normalizer (C2).
    Converts None, NaN, empty strings, whitespace-only, 'null', 'none', '<na>' into None.
    """
    if val is None or pd.isna(val):
        return None
    s = str(val).strip()
    if not s or s.lower() in ("nan", "none", "null", "<na>"):
        return None
    return s


def is_valid_key(val: Any) -> bool:
    """Validates if a primary key value is present and non-blank."""
    return normalize_text(val) is not None


def canonicalize_key(value: Any) -> Optional[str]:
    """
    Canonicalizes primary key values across CSV (string) and XLSX (numeric float).
    Example: 1001.0 (float) -> '1001', ' 1001 ' -> '1001', '00123' -> '00123'.
    """
    if not is_valid_key(value):
        return None

    if isinstance(value, float) and value.is_integer():
        return str(int(value))

    return str(value).strip()


def read_spreadsheet(file_path: str, source_name: str = "Source") -> pd.DataFrame:
    """
    Reads a CSV or XLSX file into a pandas DataFrame (C1).
    Preserves raw leading zeros and string identifiers:
    - CSV: dtype=str, keep_default_na=False
    - XLSX: dtype=object, keep_default_na=False, engine="openpyxl"
    """
    file_path = file_path.strip().strip('"').strip("'")
    if not os.path.exists(file_path):
        raise IngestionError(f"{source_name} file not found: '{file_path}'")

    _, ext = os.path.splitext(file_path)
    ext = ext.lower()

    if ext not in SUPPORTED_EXTENSIONS:
        raise IngestionError(f"Unsupported file extension '{ext}' for {source_name}. Supported types: {', '.join(sorted(SUPPORTED_EXTENSIONS))}")

    try:
        if ext == ".csv":
            df = pd.read_csv(file_path, dtype=str, keep_default_na=False)
        else:
            df = pd.read_excel(file_path, engine="openpyxl", dtype=object, keep_default_na=False)
    except Exception as e:
        raise IngestionError(f"Failed to read {source_name} file '{file_path}': {e}") from e

    # Reserved internal namespace check
    for col in df.columns:
        if str(col).startswith(INTERNAL_PREFIX):
            raise SchemaError(f"Input dataset {source_name} contains reserved internal column prefix '{INTERNAL_PREFIX}': '{col}'")

    row_col_name = f"{INTERNAL_PREFIX}source_row_{source_name.lower().replace(' ', '_')}"
    df[row_col_name] = range(2, len(df) + 2)

    return df


def sanitize_column_name(col: Any) -> str:
    """Normalizes column headers to lowercase snake_case."""
    s = str(col).strip().lower()
    s = re.sub(r"[^\w]+", "_", s)
    s = re.sub(r"_+", "_", s)
    s = s.strip("_")
    return s


def preprocess_data(df: pd.DataFrame, source_name: str = "Source") -> pd.DataFrame:
    """
    Standardizes DataFrame headers and trims string cells.
    Detects header collision errors.
    """
    df = df.copy()

    sanitized_map: Dict[str, List[str]] = {}
    new_columns = []

    for orig_col in df.columns:
        if str(orig_col).startswith(INTERNAL_PREFIX):
            new_columns.append(orig_col)
            continue
        clean_col = sanitize_column_name(orig_col)
        new_columns.append(clean_col)
        sanitized_map.setdefault(clean_col, []).append(str(orig_col))

    collisions = {k: v for k, v in sanitized_map.items() if len(v) > 1}
    if collisions:
        collision_details = "; ".join([f"'{k}': {v}" for k, v in collisions.items()])
        raise SchemaError(f"Header normalization collision detected in {source_name}: {collision_details}")

    df.columns = new_columns

    biz_cols = [c for c in df.columns if not c.startswith(INTERNAL_PREFIX)]
    df = df.dropna(how="all", subset=biz_cols)

    # Trim leading/trailing whitespace without dropping raw text strings
    df = df.map(lambda x: x.strip() if isinstance(x, str) else x)

    return df


def parse_decimal(val: Any) -> Union[Decimal, None, str]:
    """
    Safely parses financial amounts into decimal.Decimal.

    Supports:
    - 1500, 1500.00, 1,500.00, $1,500.00, ₹1,500.00, €1,500.00, USD 1500.00
    - Negative amounts: -1500.00, (500.00), ($1,500.00)
    - Indian grouping: 1,23,456.78
    """
    if val is None or pd.isna(val):
        return None

    if isinstance(val, Decimal):
        return val

    if isinstance(val, (int, float)):
        if pd.isna(val):
            return None
        try:
            return Decimal(str(val))
        except InvalidOperation:
            return INVALID_VALUE_SENTINEL

    val_str = str(val).strip()
    if not val_str or val_str.lower() in ("nan", "none", "null", "<na>"):
        return None

    is_parenthesized_negative = False
    if val_str.startswith("(") and val_str.endswith(")"):
        is_parenthesized_negative = True
        val_str = val_str[1:-1].strip()

    without_currency = CURRENCY_PATTERN.sub("", val_str).strip()

    valid_syntax = (
        re.match(r"^-?\d{1,3}(,\d{3})+(\.\d+)?$", without_currency)
        or re.match(r"^-?\d{1,3}(,\d{2})*(,\d{3})+(\.\d+)?$", without_currency)
        or re.match(r"^-?\d+(\.\d+)?$", without_currency)
    )

    if not valid_syntax:
        return INVALID_VALUE_SENTINEL

    cleaned = without_currency.replace(",", "")
    if is_parenthesized_negative and not cleaned.startswith("-"):
        cleaned = f"-{cleaned}"

    try:
        return Decimal(cleaned)
    except InvalidOperation:
        return INVALID_VALUE_SENTINEL


def parse_money(val: Any) -> ParsedMoney:
    """
    Structured monetary parser (C3).
    Extracts amount Decimal, explicit ISO currency code, symbol, and validity.
    """
    if val is None or pd.isna(val):
        return ParsedMoney(None, None, None, True)

    if isinstance(val, (int, float, Decimal)):
        dec = parse_decimal(val)
        if dec == INVALID_VALUE_SENTINEL:
            return ParsedMoney(None, None, None, False)
        return ParsedMoney(dec, None, None, True)

    val_str = str(val).strip()
    if not val_str or val_str.lower() in ("nan", "none", "null", "<na>"):
        return ParsedMoney(None, None, None, True)

    iso_match = re.search(r"\b(USD|EUR|INR|GBP|CAD|AUD|JPY|CHF|CNY|HKD|NZD)\b", val_str, re.IGNORECASE)
    explicit_curr = iso_match.group(1).upper() if iso_match else None

    sym_match = re.search(r"[$₹€£¥]|Rs\.?", val_str)
    symbol = sym_match.group(0) if sym_match else None

    if not explicit_curr and symbol in SYMBOL_TO_CURRENCY:
        explicit_curr = SYMBOL_TO_CURRENCY[symbol]

    dec = parse_decimal(val_str)
    if dec == INVALID_VALUE_SENTINEL:
        return ParsedMoney(None, explicit_curr, symbol, False)

    return ParsedMoney(dec, explicit_curr, symbol, True)


def parse_date(val: Any, fmt: Optional[str] = "%Y-%m-%d") -> Union[pd.Timestamp, None, str]:
    """Safely parses dates into normalized pd.Timestamp."""
    if val is None or pd.isna(val):
        return None

    if isinstance(val, (pd.Timestamp, datetime, date)):
        return pd.Timestamp(val).normalize()

    val_str = str(val).strip()
    if not val_str or val_str.lower() in ("nan", "none", "null", "<na>"):
        return None

    try:
        if fmt:
            parsed = pd.to_datetime(val_str, format=fmt, errors="coerce")
        else:
            parsed = pd.to_datetime(val_str, errors="coerce")

        if pd.isna(parsed):
            return INVALID_VALUE_SENTINEL
        return parsed.normalize()
    except Exception:
        return INVALID_VALUE_SENTINEL


def summarize_money_values(series: pd.Series) -> Tuple[Decimal, int, int]:
    """
    Safely sums monetary values without converting invalid amounts to Decimal("0") (C4).
    Returns (valid_total_decimal, invalid_count, missing_count).
    """
    total = Decimal("0")
    invalid_count = 0
    missing_count = 0

    for raw_value in series:
        parsed = parse_money(raw_value)
        if isinstance(parsed.amount, Decimal):
            total += parsed.amount
        elif not parsed.valid:
            invalid_count += 1
        elif parsed.amount is None:
            missing_count += 1

    return total, invalid_count, missing_count


def sanitize_excel_cell_value(val: Any) -> Any:
    """Prevents Excel formula injection by prefixing text starting with '=', '+', '-', '@' with a single quote."""
    if isinstance(val, str) and val:
        if val[0] in ("=", "+", "-", "@"):
            if val[0] == "-" and re.match(r"^-?\d+(\.\d+)?$", val):
                return val
            return f"'{val}"
    return val


def get_default_field_rules(tolerance: Decimal) -> Dict[str, Dict[str, Any]]:
    """Returns explicit default field rules."""
    return {
        "amount": {
            "type": "money",
            "tolerance": tolerance,
            "required": True,
            "currency_field": "currency",
        },
        "currency": {
            "type": "text",
            "case_sensitive": False,
            "required": True,
        },
        "customer_name": {
            "type": "text",
            "case_sensitive": False,
            "required": False,
        },
        "date": {
            "type": "date",
            "format": "%Y-%m-%d",
            "required": True,
        },
    }


def normalize_and_validate_field_rules(field_rules: Dict[str, Dict[str, Any]], key_column: str) -> Dict[str, Dict[str, Any]]:
    """
    Normalizes configured field names to snake_case exactly once.
    Validates rule types, tolerances, required parameters, and validity.
    """
    key_norm = sanitize_column_name(key_column)
    normalized_rules: Dict[str, Dict[str, Any]] = {}

    for orig_name, rule in field_rules.items():
        if not isinstance(rule, dict):
            raise ConfigurationError(f"Field rule for '{orig_name}' must be a dictionary.")

        norm_name = sanitize_column_name(orig_name)
        if norm_name == key_norm:
            raise ConfigurationError(f"Primary key column '{key_column}' cannot be included in comparison field_rules.")
        if norm_name in normalized_rules:
            raise ConfigurationError(f"Duplicate field rule after normalization: '{orig_name}' and another field both normalize to '{norm_name}'")

        f_type = str(rule.get("type", "")).strip().lower()
        if f_type not in ("money", "text", "date"):
            raise ConfigurationError(f"Invalid field type '{rule.get('type')}' for field '{orig_name}'. Allowed types: 'money', 'text', 'date'.")

        req = rule.get("required", False)
        if not isinstance(req, bool):
            raise ConfigurationError(f"Parameter 'required' for field '{orig_name}' must be a boolean.")

        validated_rule: Dict[str, Any] = {"type": f_type, "required": req}

        if f_type == "money":
            raw_tol = rule.get("tolerance", DEFAULT_TOLERANCE)
            try:
                tol_dec = Decimal(str(raw_tol))
            except (InvalidOperation, ValueError, TypeError):
                raise ConfigurationError(f"Invalid monetary tolerance '{raw_tol}' for field '{orig_name}'.")
            if tol_dec < Decimal("0"):
                raise ConfigurationError(f"Monetary tolerance for field '{orig_name}' cannot be negative: {tol_dec}")
            validated_rule["tolerance"] = tol_dec

            curr_field = rule.get("currency_field", None)
            if curr_field is not None:
                validated_rule["currency_field"] = sanitize_column_name(curr_field)

        elif f_type == "text":
            case_sens = rule.get("case_sensitive", False)
            if not isinstance(case_sens, bool):
                raise ConfigurationError(f"Parameter 'case_sensitive' for field '{orig_name}' must be a boolean.")
            validated_rule["case_sensitive"] = case_sens

        elif f_type == "date":
            fmt = rule.get("format", "%Y-%m-%d")
            if fmt is not None and not isinstance(fmt, str):
                raise ConfigurationError(f"Parameter 'format' for field '{orig_name}' must be a string format.")
            validated_rule["format"] = fmt

        normalized_rules[norm_name] = validated_rule

    if not normalized_rules:
        raise ConfigurationError("Zero valid comparison field rules were provided.")

    return normalized_rules


def load_config_file(config_path: str) -> Tuple[Dict[str, Any], str, str]:
    """Loads configuration JSON file and returns (config_dict, key_column, config_hash)."""
    config_path = config_path.strip().strip('"').strip("'")
    if not os.path.exists(config_path):
        raise ConfigurationError(f"Configuration file not found: '{config_path}'")

    config_hash = calculate_file_hash(config_path)
    try:
        with open(config_path, "r", encoding="utf-8") as f:
            data = json.load(f)
    except Exception as e:
        raise ConfigurationError(f"Failed to parse JSON config file '{config_path}': {e}") from e

    return data, data.get("key_column", DEFAULT_KEY_COLUMN), config_hash


def reconcile_data(
    df_erp: pd.DataFrame,
    df_external: pd.DataFrame,
    key_column: str = DEFAULT_KEY_COLUMN,
    amount_tolerance: Union[float, Decimal] = DEFAULT_TOLERANCE,
    field_rules: Optional[Dict[str, Dict[str, Any]]] = None,
    source_erp_name: str = "ERP",
    source_external_name: str = "External",
) -> Dict[str, pd.DataFrame]:
    """
    Core Financial Reconciliation Engine.
    """
    run_id = str(uuid.uuid4())
    run_timestamp_utc = datetime.now(timezone.utc).isoformat()

    if isinstance(amount_tolerance, (float, int, str)):
        try:
            amount_tolerance = Decimal(str(amount_tolerance))
        except InvalidOperation:
            raise ConfigurationError(f"Invalid amount tolerance value: {amount_tolerance}")

    if amount_tolerance < Decimal("0"):
        raise ConfigurationError(f"Amount tolerance cannot be negative: {amount_tolerance}")

    key_col = sanitize_column_name(key_column)

    # 1. Key Column Presence Validation
    if key_col not in df_erp.columns:
        raise SchemaError(
            f"Key column '{key_col}' not found in {source_erp_name}. Available columns: {[c for c in df_erp.columns if not c.startswith(INTERNAL_PREFIX)]}"
        )
    if key_col not in df_external.columns:
        raise SchemaError(
            f"Key column '{key_col}' not found in {source_external_name}. Available columns: {[c for c in df_external.columns if not c.startswith(INTERNAL_PREFIX)]}"
        )

    # 2. Setup Field Rules & Enforce Mandatory Schema Validation (C2)
    if field_rules is None:
        field_rules = get_default_field_rules(amount_tolerance)

    normalized_rules = normalize_and_validate_field_rules(field_rules, key_col)

    # Mandatory Schema Check: Ensure every configured rule field exists in both datasets (no silent shrinking)
    for clean_c in normalized_rules.keys():
        missing_in = []
        if clean_c not in df_erp.columns:
            missing_in.append(source_erp_name)
        if clean_c not in df_external.columns:
            missing_in.append(source_external_name)
        if missing_in:
            raise SchemaError(f"Required comparison field '{clean_c}' is missing in: {', '.join(missing_in)}")

    comparison_fields = list(normalized_rules.keys())

    logger.info(f"Starting reconciliation run {run_id}. Key: '{key_col}'. Compared fields: {comparison_fields}")

    # 3. Canonicalize Reconciliation Key (_recon_key)
    df_erp[RECON_KEY] = df_erp[key_col].map(canonicalize_key)
    df_external[RECON_KEY] = df_external[key_col].map(canonicalize_key)

    # 4. Primary Key Null / Invalid Quarantine
    if df_erp.empty:
        df_erp_valid = df_erp.copy()
        df_erp_invalid_keys = df_erp.copy()
    else:
        erp_valid_key_mask = df_erp[RECON_KEY].notna()
        df_erp_invalid_keys = df_erp[~erp_valid_key_mask].copy()
        df_erp_valid = df_erp[erp_valid_key_mask].copy()

    if df_external.empty:
        df_ext_valid = df_external.copy()
        df_ext_invalid_keys = df_external.copy()
    else:
        ext_valid_key_mask = df_external[RECON_KEY].notna()
        df_ext_invalid_keys = df_external[~ext_valid_key_mask].copy()
        df_ext_valid = df_external[ext_valid_key_mask].copy()

    invalid_key_records = []
    if not df_erp_invalid_keys.empty:
        df_erp_invalid_keys["reconciliation_status"] = STATUS_MISSING_IN_EXTERNAL
        df_erp_invalid_keys["data_quality_status"] = DQ_MISSING_PRIMARY_KEY
        df_erp_invalid_keys["status_reason"] = f"Missing or blank primary key in {source_erp_name}"
        df_erp_invalid_keys["source_system"] = source_erp_name
        invalid_key_records.append(df_erp_invalid_keys)

    if not df_ext_invalid_keys.empty:
        df_ext_invalid_keys["reconciliation_status"] = STATUS_MISSING_IN_ERP
        df_ext_invalid_keys["data_quality_status"] = DQ_MISSING_PRIMARY_KEY
        df_ext_invalid_keys["status_reason"] = f"Missing or blank primary key in {source_external_name}"
        df_ext_invalid_keys["source_system"] = source_external_name
        invalid_key_records.append(df_ext_invalid_keys)

    # 5. Deterministic Duplicate Quarantine on RECON_KEY
    if df_erp_valid.empty or RECON_KEY not in df_erp_valid.columns:
        erp_counts = pd.Series(dtype=int)
    else:
        erp_counts = df_erp_valid[RECON_KEY].value_counts()

    if df_ext_valid.empty or RECON_KEY not in df_ext_valid.columns:
        ext_counts = pd.Series(dtype=int)
    else:
        ext_counts = df_ext_valid[RECON_KEY].value_counts()

    erp_dupe_keys = set(erp_counts[erp_counts > 1].index)
    ext_dupe_keys = set(ext_counts[ext_counts > 1].index)
    all_dupe_keys_set = erp_dupe_keys.union(ext_dupe_keys)

    sorted_dupe_keys = sorted(list(all_dupe_keys_set), key=str)

    if df_erp_valid.empty or RECON_KEY not in df_erp_valid.columns:
        df_erp_dupes = pd.DataFrame(columns=df_erp.columns)
        df_erp_clean = df_erp_valid.copy()
    else:
        df_erp_dupes = df_erp_valid[df_erp_valid[RECON_KEY].isin(all_dupe_keys_set)].copy()
        df_erp_clean = df_erp_valid[~df_erp_valid[RECON_KEY].isin(all_dupe_keys_set)].copy()

    if df_ext_valid.empty or RECON_KEY not in df_ext_valid.columns:
        df_ext_dupes = pd.DataFrame(columns=df_external.columns)
        df_ext_clean = df_ext_valid.copy()
    else:
        df_ext_dupes = df_ext_valid[df_ext_valid[RECON_KEY].isin(all_dupe_keys_set)].copy()
        df_ext_clean = df_ext_valid[~df_ext_valid[RECON_KEY].isin(all_dupe_keys_set)].copy()

    dupe_records = []
    for k in sorted_dupe_keys:
        in_erp = k in erp_dupe_keys
        in_ext = k in ext_dupe_keys

        if in_erp and in_ext:
            dq_status = DQ_DUPLICATE_IN_BOTH
            reason = f"Duplicate primary key '{k}' detected in both {source_erp_name} and {source_external_name}"
        elif in_erp:
            dq_status = DQ_DUPLICATE_IN_ERP
            reason = f"Duplicate primary key '{k}' detected in {source_erp_name}"
        else:
            dq_status = DQ_DUPLICATE_IN_EXTERNAL
            reason = f"Duplicate primary key '{k}' detected in {source_external_name}"

        rows_erp = df_erp_dupes[df_erp_dupes[RECON_KEY] == k].copy()
        if not rows_erp.empty:
            rows_erp["reconciliation_status"] = STATUS_MISSING_IN_EXTERNAL
            rows_erp["data_quality_status"] = dq_status
            rows_erp["status_reason"] = reason
            rows_erp["source_system"] = source_erp_name
            dupe_records.append(rows_erp)

        rows_ext = df_ext_dupes[df_ext_dupes[RECON_KEY] == k].copy()
        if not rows_ext.empty:
            rows_ext["reconciliation_status"] = STATUS_MISSING_IN_ERP
            rows_ext["data_quality_status"] = dq_status
            rows_ext["status_reason"] = reason
            rows_ext["source_system"] = source_external_name
            dupe_records.append(rows_ext)

    # 6. One-to-One Vector Outer Merge on RECON_KEY
    merged = df_erp_clean.merge(
        df_ext_clean,
        on=RECON_KEY,
        how="outer",
        suffixes=(f"_{source_erp_name.lower()}", f"_{source_external_name.lower()}"),
        indicator="_merge",
        validate="one_to_one",
    )

    erp_suffix = f"_{source_erp_name.lower()}"
    ext_suffix = f"_{source_external_name.lower()}"

    missing_records = merged[merged["_merge"] != "both"].copy()
    missing_records["reconciliation_status"] = missing_records["_merge"].map(
        {
            "left_only": STATUS_MISSING_IN_EXTERNAL,
            "right_only": STATUS_MISSING_IN_ERP,
        }
    )
    missing_records["data_quality_status"] = DQ_VALID
    missing_records["status_reason"] = missing_records["_merge"].map(
        {
            "left_only": f"Present in {source_erp_name} only (Missing in {source_external_name})",
            "right_only": f"Present in {source_external_name} only (Missing in {source_erp_name})",
        }
    )
    missing_records["source_system"] = missing_records["_merge"].map(
        {
            "left_only": source_erp_name,
            "right_only": source_external_name,
        }
    )

    # Perform Source-Level Data Quality checks on Missing records (C8)
    for idx, row in missing_records.iterrows():
        is_left = row["_merge"] == "left_only"
        suff = erp_suffix if is_left else ext_suffix
        src_name = source_erp_name if is_left else source_external_name

        for col in comparison_fields:
            rule = normalized_rules[col]
            val = row.get(f"{col}{suff}")
            f_type = rule["type"]
            is_req = rule.get("required", False)

            if f_type == "money":
                pm = parse_money(val)
                if not pm.valid:
                    missing_records.at[idx, "data_quality_status"] = DQ_INVALID_FIELD_VALUE
                    missing_records.at[idx, "status_reason"] += f"; {col} has invalid monetary format in {src_name}"
                elif pm.amount is None and is_req:
                    missing_records.at[idx, "data_quality_status"] = DQ_MISSING_REQUIRED_VALUE
                    missing_records.at[idx, "status_reason"] += f"; {col} is required but blank in {src_name}"

            elif f_type == "text":
                if normalize_text(val) is None and is_req:
                    missing_records.at[idx, "data_quality_status"] = DQ_MISSING_REQUIRED_VALUE
                    missing_records.at[idx, "status_reason"] += f"; {col} is required but blank in {src_name}"

            elif f_type == "date":
                fmt = rule.get("format", "%Y-%m-%d")
                parsed_dt = parse_date(val, fmt=fmt)
                if parsed_dt == INVALID_VALUE_SENTINEL:
                    missing_records.at[idx, "data_quality_status"] = DQ_INVALID_FIELD_VALUE
                    missing_records.at[idx, "status_reason"] += f"; {col} has invalid date format in {src_name}"
                elif parsed_dt is None and is_req:
                    missing_records.at[idx, "data_quality_status"] = DQ_MISSING_REQUIRED_VALUE
                    missing_records.at[idx, "status_reason"] += f"; {col} is required but blank in {src_name}"

    # 7. Field-Level Typed Comparison for Common Records
    common = merged[merged["_merge"] == "both"].copy()

    common["reconciliation_status"] = STATUS_MATCH
    common["data_quality_status"] = DQ_VALID
    common["status_reason"] = ""
    common["mismatch_fields"] = ""
    common["mismatch_count"] = 0

    for col in comparison_fields:
        rule = normalized_rules[col]
        if rule["type"] == "money":
            common[f"{col}_variance"] = None
            common[f"{col}_abs_difference"] = None

    recon_statuses = []
    dq_statuses = []
    status_reasons = []
    mismatch_fields_list = []
    mismatch_counts = []
    currency_mismatch_flags = []

    for idx, row in common.iterrows():
        row_mismatched_fields = []
        row_reasons = []
        row_has_invalid = False
        row_has_missing_required = False
        row_has_curr_conflict = False
        is_currency_mismatch = False

        for col in comparison_fields:
            rule = normalized_rules[col]
            f_type = rule["type"]
            is_req = rule.get("required", False)

            val_erp_raw = row.get(f"{col}{erp_suffix}")
            val_ext_raw = row.get(f"{col}{ext_suffix}")

            if f_type == "money":
                col_tol = rule.get("tolerance", amount_tolerance)
                pm_erp = parse_money(val_erp_raw)
                pm_ext = parse_money(val_ext_raw)

                if not pm_erp.valid or not pm_ext.valid:
                    row_has_invalid = True
                    row_reasons.append(f"{col} has invalid monetary format ({source_erp_name}: '{val_erp_raw}', {source_external_name}: '{val_ext_raw}')")
                else:
                    # Currency Symbol / Code Conflict Check (C3)
                    curr_col = rule.get("currency_field", "currency")
                    rec_curr_erp = normalize_text(row.get(f"{curr_col}{erp_suffix}")) if curr_col in df_erp.columns else None
                    rec_curr_ext = normalize_text(row.get(f"{curr_col}{ext_suffix}")) if curr_col in df_external.columns else None

                    if pm_erp.explicit_currency and rec_curr_erp and pm_erp.explicit_currency.upper() != rec_curr_erp.upper():
                        row_has_curr_conflict = True
                        row_reasons.append(
                            f"{col} embedded currency '{pm_erp.explicit_currency}' conflicts with currency field '{rec_curr_erp}' in {source_erp_name}"
                        )
                    if pm_ext.explicit_currency and rec_curr_ext and pm_ext.explicit_currency.upper() != rec_curr_ext.upper():
                        row_has_curr_conflict = True
                        row_reasons.append(
                            f"{col} embedded currency '{pm_ext.explicit_currency}' conflicts with currency field '{rec_curr_ext}' in {source_external_name}"
                        )

                    dec_erp = pm_erp.amount
                    dec_ext = pm_ext.amount

                    if dec_erp is None and dec_ext is None:
                        if is_req:
                            row_has_missing_required = True
                            row_reasons.append(f"{col} is required but missing in both datasets")
                            common.at[idx, f"{col}_variance"] = None
                            common.at[idx, f"{col}_abs_difference"] = None
                        else:
                            common.at[idx, f"{col}_variance"] = Decimal("0")
                            common.at[idx, f"{col}_abs_difference"] = Decimal("0")
                    elif dec_erp is None or dec_ext is None:
                        if is_req:
                            row_has_missing_required = True
                        row_mismatched_fields.append(col)
                        missing_where = source_erp_name if dec_erp is None else source_external_name
                        row_reasons.append(f"{col} is required but missing in {missing_where}")
                    else:
                        variance = dec_ext - dec_erp  # External - ERP
                        abs_diff = abs(variance)
                        common.at[idx, f"{col}_variance"] = variance
                        common.at[idx, f"{col}_abs_difference"] = abs_diff

                        if abs_diff > col_tol:
                            row_mismatched_fields.append(col)
                            row_reasons.append(f"{col} diff ({source_erp_name}: {dec_erp}, {source_external_name}: {dec_ext}, variance: {variance})")

            elif f_type == "text":
                case_sens = rule.get("case_sensitive", False)
                str_erp = normalize_text(val_erp_raw)
                str_ext = normalize_text(val_ext_raw)

                if str_erp is None and str_ext is None:
                    if is_req:
                        row_has_missing_required = True
                        row_reasons.append(f"{col} is required but missing in both datasets")
                elif str_erp is None or str_ext is None:
                    if is_req:
                        row_has_missing_required = True
                    row_mismatched_fields.append(col)
                    missing_where = source_erp_name if str_erp is None else source_external_name
                    row_reasons.append(f"{col} is required but missing in {missing_where}")
                else:
                    c_erp = str_erp if case_sens else str_erp.lower()
                    c_ext = str_ext if case_sens else str_ext.lower()
                    if c_erp != c_ext:
                        row_mismatched_fields.append(col)
                        if col == "currency":
                            is_currency_mismatch = True
                        row_reasons.append(f"{col} mismatch ({source_erp_name}: '{val_erp_raw}', {source_external_name}: '{val_ext_raw}')")

            elif f_type == "date":
                fmt = rule.get("format", "%Y-%m-%d")
                dt_erp = parse_date(val_erp_raw, fmt=fmt)
                dt_ext = parse_date(val_ext_raw, fmt=fmt)

                if dt_erp == INVALID_VALUE_SENTINEL or dt_ext == INVALID_VALUE_SENTINEL:
                    row_has_invalid = True
                    row_reasons.append(f"{col} has unparseable date format ({source_erp_name}: '{val_erp_raw}', {source_external_name}: '{val_ext_raw}')")
                elif dt_erp is None and dt_ext is None:
                    if is_req:
                        row_has_missing_required = True
                        row_reasons.append(f"{col} is required but missing in both datasets")
                elif dt_erp is None or dt_ext is None:
                    if is_req:
                        row_has_missing_required = True
                    row_mismatched_fields.append(col)
                    missing_where = source_erp_name if dt_erp is None else source_external_name
                    row_reasons.append(f"{col} date missing in {missing_where}")
                elif dt_erp != dt_ext:
                    row_mismatched_fields.append(col)
                    row_reasons.append(
                        f"{col} date mismatch ({source_erp_name}: {dt_erp.strftime('%Y-%m-%d')}, {source_external_name}: {dt_ext.strftime('%Y-%m-%d')})"
                    )

        # Assign reconciliation_status and data_quality_status independently (C8)
        if row_mismatched_fields or row_has_missing_required or row_has_invalid or row_has_curr_conflict:
            r_status = STATUS_MISMATCH
        else:
            r_status = STATUS_MATCH

        if row_has_invalid:
            dq_status = DQ_INVALID_FIELD_VALUE
        elif row_has_curr_conflict:
            dq_status = DQ_CURRENCY_AMOUNT_CONFLICT
        elif row_has_missing_required:
            dq_status = DQ_MISSING_REQUIRED_VALUE
        else:
            dq_status = DQ_VALID

        recon_statuses.append(r_status)
        dq_statuses.append(dq_status)
        status_reasons.append("; ".join(row_reasons))
        mismatch_fields_list.append(", ".join(row_mismatched_fields))
        mismatch_counts.append(len(row_mismatched_fields))
        currency_mismatch_flags.append(is_currency_mismatch)

    common["reconciliation_status"] = recon_statuses
    common["data_quality_status"] = dq_statuses
    common["status_reason"] = status_reasons
    common["mismatch_fields"] = mismatch_fields_list
    common["mismatch_count"] = mismatch_counts
    common["is_currency_mismatch"] = currency_mismatch_flags

    # 8. Aggregate Unified Category DataFrames
    df_matches = common[common["reconciliation_status"] == STATUS_MATCH].copy()
    df_mismatches = common[common["reconciliation_status"] == STATUS_MISMATCH].copy()

    # Source Exceptions Grain (C4): 1 row per physical problematic source record
    source_exception_components = []
    if invalid_key_records:
        source_exception_components.extend(invalid_key_records)
    if dupe_records:
        source_exception_components.extend(dupe_records)

    if source_exception_components:
        df_source_exceptions = pd.concat(source_exception_components, ignore_index=True)
    else:
        df_source_exceptions = pd.DataFrame(columns=[key_col, "reconciliation_status", "data_quality_status", "status_reason", "source_system"])

    df_missing = missing_records.copy()

    # 9. Build Reconciliation Units Sheet (C4): 1 row per logical reconciliation unit
    df_recon_units = pd.concat([common, df_missing], ignore_index=True)
    df_recon_units = df_recon_units.drop(columns=["_merge"], errors="ignore")

    user_key_cols = [c for c in [key_col, f"{key_col}_{source_erp_name.lower()}", f"{key_col}_{source_external_name.lower()}"] if c in df_recon_units.columns]
    primary_cols = user_key_cols + ["reconciliation_status", "data_quality_status", "status_reason", "mismatch_fields", "mismatch_count"]
    other_cols = [c for c in df_recon_units.columns if c not in primary_cols and c != RECON_KEY]
    df_recon_units = df_recon_units[primary_cols + other_cols]

    # 10. Compute Explicit KPIs and Dynamic Currency-Isolated Exposure (C5 & C7)
    total_erp_ingested = len(df_erp)
    total_ext_ingested = len(df_external)
    total_source_rows = total_erp_ingested + total_ext_ingested

    total_source_exception_rows = len(df_source_exceptions)
    source_row_exception_rate = (total_source_exception_rows / total_source_rows * 100) if total_source_rows > 0 else 0.0

    eligible_count = len(common)
    matches_count = len(df_matches)
    mismatches_count = len(df_mismatches)
    missing_ext_count = len(df_missing[df_missing["reconciliation_status"] == STATUS_MISSING_IN_EXTERNAL])
    missing_erp_count = len(df_missing[df_missing["reconciliation_status"] == STATUS_MISSING_IN_ERP])

    total_recon_units = len(df_recon_units)
    unreconciled_units_count = total_recon_units - matches_count
    recon_unit_exception_rate = (unreconciled_units_count / total_recon_units * 100) if total_recon_units > 0 else 0.0
    eligible_match_rate = (matches_count / eligible_count * 100) if eligible_count > 0 else 0.0

    distinct_issue_keys_count = (
        df_source_exceptions[RECON_KEY].nunique() if (not df_source_exceptions.empty and RECON_KEY in df_source_exceptions.columns) else 0
    )
    currency_mismatch_count = common["is_currency_mismatch"].sum() if "is_currency_mismatch" in common.columns else 0

    summary_rows = [
        {"Metric / Indicator": "Total ERP Source Rows Ingested", "Value": total_erp_ingested},
        {"Metric / Indicator": "Total External Source Rows Ingested", "Value": total_ext_ingested},
        {"Metric / Indicator": "Total Ingested Source Rows", "Value": total_source_rows},
        {"Metric / Indicator": "Total Source Exception Rows (Data Quality)", "Value": total_source_exception_rows},
        {"Metric / Indicator": "Source Row Exception Rate (%)", "Value": f"{source_row_exception_rate:.2f}%"},
        {"Metric / Indicator": "Distinct Data Quality Exception Keys", "Value": distinct_issue_keys_count},
        {"Metric / Indicator": "Eligible 1-to-1 Reconciliation Units", "Value": eligible_count},
        {"Metric / Indicator": "Fully Matched Reconciliation Units (MATCH)", "Value": matches_count},
        {"Metric / Indicator": "Mismatch Reconciliation Units (MISMATCH)", "Value": mismatches_count},
        {"Metric / Indicator": "Missing in External Units (MISSING_IN_EXTERNAL)", "Value": missing_ext_count},
        {"Metric / Indicator": "Missing in ERP Units (MISSING_IN_ERP)", "Value": missing_erp_count},
        {"Metric / Indicator": "Currency Mismatch Count", "Value": int(currency_mismatch_count)},
        {"Metric / Indicator": "Eligible Match Rate (%)", "Value": f"{eligible_match_rate:.2f}%"},
        {"Metric / Indicator": "Reconciliation Unit Exception Rate (%)", "Value": f"{recon_unit_exception_rate:.2f}%"},
    ]

    # Calculate monetary exposure dynamically for each configured money field (C7)
    for col, rule in normalized_rules.items():
        if rule["type"] == "money":
            curr_col = rule.get("currency_field", "currency")
            curr_erp_col = f"{curr_col}_{source_erp_name.lower()}"
            curr_ext_col = f"{curr_col}_{source_external_name.lower()}"

            # Mismatched Exposure
            if not df_mismatches.empty and f"{col}_abs_difference" in df_mismatches.columns:
                if curr_erp_col in df_mismatches.columns and curr_ext_col in df_mismatches.columns:
                    same_curr_df = df_mismatches[df_mismatches[curr_erp_col].astype(str).str.lower() == df_mismatches[curr_ext_col].astype(str).str.lower()]
                    for curr in same_curr_df[curr_erp_col].dropna().unique():
                        c_df = same_curr_df[same_curr_df[curr_erp_col] == curr]
                        exposure = c_df[f"{col}_abs_difference"].dropna().sum()
                        summary_rows.append(
                            {
                                "Metric / Indicator": f"Mismatched {col.title()} Exposure ({curr})",
                                "Value": f"{exposure:.2f}",
                            }
                        )

            # Missing Exposure (C4 & C7)
            if not df_missing.empty:
                amt_erp_col = f"{col}_{source_erp_name.lower()}"
                amt_ext_col = f"{col}_{source_external_name.lower()}"

                missing_ext_df = df_missing[df_missing["reconciliation_status"] == STATUS_MISSING_IN_EXTERNAL]
                if not missing_ext_df.empty and amt_erp_col in missing_ext_df.columns and curr_erp_col in missing_ext_df.columns:
                    for curr in missing_ext_df[curr_erp_col].dropna().unique():
                        c_df = missing_ext_df[missing_ext_df[curr_erp_col] == curr]
                        total_amt, inv_cnt, blank_cnt = summarize_money_values(c_df[amt_erp_col])
                        summary_rows.append(
                            {
                                "Metric / Indicator": f"Missing in External {col.title()} Exposure ({curr})",
                                "Value": f"{total_amt:.2f}",
                            }
                        )
                        if inv_cnt > 0:
                            summary_rows.append(
                                {
                                    "Metric / Indicator": f"Missing in External Rows With Invalid {col.title()} ({curr})",
                                    "Value": inv_cnt,
                                }
                            )
                        if blank_cnt > 0:
                            summary_rows.append(
                                {
                                    "Metric / Indicator": f"Missing in External Rows With Blank {col.title()} ({curr})",
                                    "Value": blank_cnt,
                                }
                            )

                missing_erp_df = df_missing[df_missing["reconciliation_status"] == STATUS_MISSING_IN_ERP]
                if not missing_erp_df.empty and amt_ext_col in missing_erp_df.columns and curr_ext_col in missing_erp_df.columns:
                    for curr in missing_erp_df[curr_ext_col].dropna().unique():
                        c_df = missing_erp_df[missing_erp_df[curr_ext_col] == curr]
                        total_amt, inv_cnt, blank_cnt = summarize_money_values(c_df[amt_ext_col])
                        summary_rows.append(
                            {
                                "Metric / Indicator": f"Missing in ERP {col.title()} Exposure ({curr})",
                                "Value": f"{total_amt:.2f}",
                            }
                        )
                        if inv_cnt > 0:
                            summary_rows.append(
                                {
                                    "Metric / Indicator": f"Missing in ERP Rows With Invalid {col.title()} ({curr})",
                                    "Value": inv_cnt,
                                }
                            )
                        if blank_cnt > 0:
                            summary_rows.append(
                                {
                                    "Metric / Indicator": f"Missing in ERP Rows With Blank {col.title()} ({curr})",
                                    "Value": blank_cnt,
                                }
                            )

    df_summary = pd.DataFrame(summary_rows)

    # 11. Run Metadata Sheet
    serialized_rules = json.dumps(
        {k: {rk: str(rv) for rk, rv in v.items()} for k, v in normalized_rules.items()},
        indent=2,
    )

    df_metadata = pd.DataFrame(
        [
            {"Parameter": "Run ID", "Value": run_id},
            {"Parameter": "Run Timestamp UTC", "Value": run_timestamp_utc},
            {"Parameter": "Tool Version", "Value": get_tool_version()},
            {"Parameter": "Primary Key Column", "Value": key_col},
            {"Parameter": "Canonical Key Field", "Value": RECON_KEY},
            {"Parameter": "Amount Tolerance Threshold", "Value": str(amount_tolerance)},
            {"Parameter": "Compared Fields", "Value": ", ".join(comparison_fields)},
            {"Parameter": "Serialized Field Rules", "Value": serialized_rules},
            {"Parameter": f"{source_erp_name} Total Ingested Rows", "Value": total_erp_ingested},
            {"Parameter": f"{source_external_name} Total Ingested Rows", "Value": total_ext_ingested},
            {"Parameter": "Eligible 1-to-1 Reconciliation Units", "Value": eligible_count},
        ]
    )

    results = {
        "summary": df_summary,
        "run_metadata": df_metadata,
        "matches": df_matches.drop(columns=["_merge", RECON_KEY], errors="ignore"),
        "mismatches": df_mismatches.drop(columns=["_merge", RECON_KEY], errors="ignore"),
        "missing": df_missing.drop(columns=["_merge", RECON_KEY], errors="ignore"),
        "source_exceptions": df_source_exceptions.drop(columns=[RECON_KEY], errors="ignore"),
        "reconciliation_units": df_recon_units.drop(columns=[RECON_KEY], errors="ignore"),
    }

    return results


def write_output(
    results: Dict[str, pd.DataFrame],
    output_folder: str,
    output_filename: Optional[str] = None,
    metadata_info: Optional[Dict[str, str]] = None,
    overwrite: bool = False,
) -> str:
    """
    Writes reconciliation results into a multi-sheet Excel workbook.
    Prevents silent file overwrites unless overwrite=True.
    """
    output_folder = output_folder.strip().strip('"').strip("'")
    if not os.path.exists(output_folder):
        os.makedirs(output_folder, exist_ok=True)

    if not output_filename:
        timestamp_str = datetime.now(timezone.utc).strftime("%Y%m%dT%H%M%S")
        short_id = str(uuid.uuid4())[:8]
        output_filename = f"reconciliation_{timestamp_str}_{short_id}.xlsx"

    output_path = os.path.join(output_folder, output_filename)

    if os.path.exists(output_path) and not overwrite:
        raise OutputError(f"Output file '{output_path}' already exists. Use --overwrite flag to explicitly overwrite.")

    if metadata_info and "run_metadata" in results:
        meta_df = results["run_metadata"].copy()
        for k, v in metadata_info.items():
            meta_df = pd.concat([meta_df, pd.DataFrame([{"Parameter": k, "Value": str(v)}])], ignore_index=True)
        results["run_metadata"] = meta_df

    sheet_order = [
        "summary",
        "run_metadata",
        "matches",
        "mismatches",
        "missing",
        "source_exceptions",
        "reconciliation_units",
    ]

    with pd.ExcelWriter(output_path, engine="openpyxl") as writer:
        for sheet_name in sheet_order:
            if sheet_name in results:
                df = results[sheet_name].copy()

                for col in df.columns:
                    df[col] = df[col].map(lambda value: sanitize_excel_cell_value(value) if isinstance(value, str) else value)

                df.to_excel(writer, sheet_name=sheet_name, index=False)

    wb = openpyxl.load_workbook(output_path)
    header_fill = PatternFill(start_color="1F4E78", end_color="1F4E78", fill_type="solid")
    header_font = Font(name="Segoe UI", size=10, bold=True, color="FFFFFF")
    data_font = Font(name="Segoe UI", size=10)
    thin_border = Border(
        left=Side(style="thin", color="D9D9D9"),
        right=Side(style="thin", color="D9D9D9"),
        top=Side(style="thin", color="D9D9D9"),
        bottom=Side(style="thin", color="D9D9D9"),
    )

    for sheet_name in wb.sheetnames:
        ws = wb[sheet_name]
        ws.views.sheetView[0].showGridLines = True
        ws.freeze_panes = "A2"

        for cell in ws[1]:
            cell.fill = header_fill
            cell.font = header_font
            cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)

        for row in ws.iter_rows(min_row=2):
            for cell in row:
                cell.font = data_font
                cell.border = thin_border
                if isinstance(cell.value, (int, float, Decimal)):
                    cell.alignment = Alignment(horizontal="right", vertical="center")

        for col in ws.columns:
            max_len = 0
            col_letter = get_column_letter(col[0].column)
            for cell in col:
                val_str = str(cell.value or "")
                if len(val_str) > max_len:
                    max_len = len(val_str)
            ws.column_dimensions[col_letter].width = min(max(max_len + 4, 12), 50)

    wb.save(output_path)
    logger.info(f"Reconciliation workbook successfully saved to '{output_path}'")
    return output_path


def parse_cli_args(args: Optional[List[str]] = None) -> argparse.Namespace:
    """Configures CLI argument parsing."""
    parser = argparse.ArgumentParser(
        description="Enterprise Finance Reconciliation Engine CLI",
        formatter_class=argparse.ArgumentDefaultsHelpFormatter,
    )
    parser.add_argument("--erp", type=str, help="Path to ERP / System of Record dataset (CSV/XLSX)")
    parser.add_argument("--external", type=str, help="Path to External / Bank dataset (CSV/XLSX)")
    parser.add_argument("--config", type=str, help="Path to JSON reconciliation configuration file")
    parser.add_argument("--output", type=str, default="output", help="Output directory path")
    parser.add_argument("--key", type=str, default=DEFAULT_KEY_COLUMN, help="Primary key column name")
    parser.add_argument(
        "--amount-tolerance",
        type=float,
        default=0.01,
        help="Monetary comparison tolerance threshold (e.g. 0.01)",
    )
    parser.add_argument("--overwrite", action="store_true", help="Overwrite output file if it already exists")
    parser.add_argument("--interactive", action="store_true", help="Force interactive mode prompts for missing arguments")
    return parser.parse_args(args)


def main():
    """CLI Entry Point."""
    logging.basicConfig(level=logging.INFO, format="%(asctime)s [%(levelname)s] %(message)s")

    print("=" * 70)
    print(" Enterprise Finance Reconciliation Tool ")
    print("=" * 70)

    args = parse_cli_args()

    erp_path = args.erp
    external_path = args.external
    config_path = args.config
    output_folder = args.output
    key_column = args.key
    tolerance = Decimal(str(args.amount_tolerance))
    field_rules = None
    config_hash = None

    if config_path:
        config_data, key_column_from_config, config_hash = load_config_file(config_path)
        if "key_column" in config_data:
            key_column = config_data["key_column"]
        if "field_rules" in config_data:
            field_rules = config_data["field_rules"]

    if not erp_path or not external_path or args.interactive:
        erp_path = input("Enter ERP / System of Record file path: ").strip()
        external_path = input("Enter External / Bank file path: ").strip()
        output_folder_input = input(f"Enter output directory path [Default: '{output_folder}']: ").strip()
        if output_folder_input:
            output_folder = output_folder_input
        key_input = input(f"Enter Key Column name [Default: '{key_column}']: ").strip()
        if key_input:
            key_column = key_input

    try:
        logger.info("Ingesting datasets...")
        erp_hash = calculate_file_hash(erp_path)
        ext_hash = calculate_file_hash(external_path)

        df_erp_raw = read_spreadsheet(erp_path, source_name="ERP")
        df_ext_raw = read_spreadsheet(external_path, source_name="External")

        df_erp = preprocess_data(df_erp_raw, source_name="ERP")
        df_ext = preprocess_data(df_ext_raw, source_name="External")

        logger.info("Reconciling datasets...")
        results = reconcile_data(
            df_erp,
            df_ext,
            key_column=key_column,
            amount_tolerance=tolerance,
            field_rules=field_rules,
            source_erp_name="ERP",
            source_external_name="External",
        )

        metadata_info = {
            "ERP File Path": erp_path,
            "ERP File SHA-256": erp_hash,
            "External File Path": external_path,
            "External File SHA-256": ext_hash,
        }
        if config_path:
            metadata_info["Config File Path"] = config_path
            metadata_info["Config File SHA-256"] = config_hash

        logger.info("Generating formatted Excel workbook...")
        output_file = write_output(results, output_folder, metadata_info=metadata_info, overwrite=args.overwrite)

        print("\n[SUCCESS] Reconciliation Completed Successfully!")
        print("\n[SUMMARY] Executive KPI Dashboard:")
        for _, row in results["summary"].iterrows():
            print(f"   - {row['Metric / Indicator']}: {row['Value']}")
        print(f"\n[OUTPUT] Auditable Workbook Saved To:\n   {output_file}\n")

    except ReconciliationError as re_err:
        logger.error(f"Reconciliation Validation Error: {re_err}")
        print(f"\n[ERROR] Reconciliation Error: {re_err}")
        sys.exit(1)
    except Exception as e:
        logger.exception(f"Unexpected System Error: {e}")
        print(f"\n[ERROR] Critical Error: {e}")
        sys.exit(1)


if __name__ == "__main__":
    main()
